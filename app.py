import difflib
import html
import io
import os
import re
import sqlite3
import zipfile
from datetime import date

import altair as alt
import pandas as pd
import requests
import streamlit as st
from lxml import html as lh
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="재무제표 검색", layout="wide")


def load_key():
    key = os.getenv("DART_API_KEY")
    if key:
        return key
    try:
        return st.secrets["DART_API_KEY"]
    except Exception:
        return None


API_KEY = load_key()
if not API_KEY:
    st.error("DART_API_KEY가 설정되지 않았습니다. README의 설정 방법을 확인하세요.")
    st.stop()


def safe_err(e):
    """에러 메시지에서 API 키를 가림"""
    return str(e).replace(API_KEY, "****")


AUTHOR = "made by inhyeok"
st.title("📊 재무제표 검색")
st.caption(AUTHOR)

URL = "https://opendart.fss.or.kr/api/fnlttSinglAcntAll.json"
REPRT = {"사업보고서": "11011", "반기보고서": "11012", "1분기보고서": "11013", "3분기보고서": "11014"}
DIV = {"연결": "CFS", "별도": "OFS"}
IDS = {"ifrs-full_Revenue": "매출액", "dart_OperatingIncomeLoss": "영업이익",
       "ifrs-full_ProfitLossFromOperatingActivities": "영업이익",
       "ifrs-full_Liabilities": "부채총계", "ifrs-full_Equity": "자본총계"}
# 요약 지표를 어느 표에서 찾을지 한정 (다른 표의 같은 ID가 잡히는 것 방지)
SJ_OF = {"매출액": ("IS", "CIS"), "영업이익": ("IS", "CIS"),
         "부채총계": ("BS",), "자본총계": ("BS",)}
# 한국 재무제표 표준 순서: 재무상태표 → 손익 → 포괄손익 → 자본변동표 → 현금흐름표
SJ_RANK = {"BS": 0, "IS": 1, "CIS": 2, "SCE": 3, "CF": 4}

# ===== 엑셀 서식 설정 =====
FONT = "맑은 고딕"
THIN = Side(style="thin", color="999999")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="D9D9D9")
SUB = PatternFill("solid", fgColor="F2F2F2")
NUM = '#,##0;(#,##0);"-"'
CREDIT = Font(name=FONT, size=9, italic=True, color="808080")   # 제작자 표시용 작은 회색 글씨
MAJOR = ("자산", "부채", "자본", "유동자산", "비유동자산", "유동부채", "비유동부채")
AMOUNT = re.compile(r"\(?-?\d{1,3}(,\d{3})*(\.\d+)?\)?")   # 1,000 미만 숫자도 인식
SUBHEAD = re.compile(r"^(\d+\.\d+|\(\d+\)|[가-하]\.)\s*\S")
STOP = re.compile(r"(내부회계관리제도.{0,10}(감사|검토)\s*의견|외부감사\s*실시내용)")
# 반기·분기보고서 본문: 주석 다음 섹션 제목이 나오면 멈춤
BODY_STOP = re.compile(r"^\d+\s*\.\s*(재무제표\s*$|배당에\s*관한\s*사항|증권의\s*발행|기타\s*재무에\s*관한)"
                       r"|^(Ⅳ|IV)\s*\.")
SEP_BS = re.compile(r"^\d+\s*-\s*\d+\s*\.\s*재무상태표")          # 연결 다음에 나오는 별도 재무상태표
SECTION_NOTE = re.compile(r"재무제표\s*주석\s*$")                     # '3. 연결재무제표 주석' 섹션 제목
NOTE_HEAD = re.compile(r"^(?:주\s*석\s*)?(\d{1,2})\s*[\.．)]\s*(?=[^\d\s])")   # 1. / 1) / 1 . / 주석 1.
NUM_ONLY = re.compile(r"^(?:주\s*석\s*)?\d{1,2}\s*[\.．)]$")                    # "1." 만 따로 떨어진 줄
# 주석 제목 뒤에 본문이 붙어 나올 때 나눌 위치: (1) / 1) / ① / 가.  또는 본문 첫머리에 자주 오는 말
SUB_START = re.compile(r"\(\d+\)|\d+\)|[①-⑳]|(?<![가-힣])[가-하]\.\s")   # '다.'(문장 끝)는 제외
BODY_WORD = re.compile(r"당기|전기|연결회사|지배기업|보고기간|연결실체")
# 앞 문단 끝에 다음 주석 제목이 붙은 경우: "...참조).13. 유형자산" → 마침표·괄호 뒤의 'N. 한글' 앞에서 나눔
MID_HEAD = re.compile(r"(?<=[\.\)\]」’'\"])\s*(?=\d{1,2}\s*\.\s*[가-힣])")


# ===== 1. XBRL 재무제표 엑셀 =====
def to_num(v):
    v = pd.to_numeric(v, errors="coerce")
    return None if pd.isna(v) else float(v)


def amount_cols(part):
    """값이 실제로 들어 있는 금액 칸만 골라 (표시이름, 컬럼명) 목록으로 돌려줌"""
    def has(c):
        return c in part.columns and pd.to_numeric(part[c], errors="coerce").notna().any()
    cols = []
    for who, base in (("당기", "thstrm"), ("전기", "frmtrm")):
        a, b = f"{base}_amount", f"{base}_add_amount"
        if has(a) and has(b):
            cols += [(f"{who}(3개월)", a), (f"{who}(누적)", b)]
        elif has(a):
            cols.append((who, a))
        elif has(b):
            cols.append((f"{who}(누적)", b))
    return cols


def sce_frame(part, col):
    """자본변동표: 행=계정(ord 순서), 열=자본 구성요소로 펼친 표"""
    rows, cols, data, seen = [], [], {}, {}
    for _, r in part.iterrows():
        nm = str(r["account_nm"])
        d = r.get("account_detail")
        parts = [re.sub(r"\s*\[member\]", "", p).strip() for p in str(d).split("|")] if isinstance(d, str) else []
        parts = [p for p in parts if p and p != "-"]
        c = parts[-1] if parts else "합계"
        n = seen.get((nm, c), 0)
        seen[(nm, c)] = n + 1
        row = nm if n == 0 else f"{nm} ({n + 1})"
        if row not in rows:
            rows.append(row)
        if c not in cols:
            cols.append(c)
        data[(row, c)] = to_num(r.get(col))
    f = pd.DataFrame(index=rows, columns=cols, dtype=float)
    for (row, c), v in data.items():
        f.loc[row, c] = v
    return f


def write_sce(ws, part, cols, r):
    for label, col in cols:
        f = sce_frame(part, col)
        ws.cell(row=r, column=1, value=label).font = Font(name=FONT, bold=True)
        r += 1
        for c, h in enumerate(["과목"] + list(f.columns), 1):
            cell = ws.cell(row=r, column=c, value=h)
            cell.font = Font(name=FONT, bold=True)
            cell.fill, cell.border = HEAD, BOX
            cell.alignment = Alignment(horizontal="center", wrap_text=True)
        r += 1
        for nm, row in f.iterrows():
            cell = ws.cell(row=r, column=1, value=nm)
            cell.border, cell.font = BOX, Font(name=FONT)
            for c, v in enumerate(row, 2):
                cell = ws.cell(row=r, column=c, value=None if pd.isna(v) else float(v))
                cell.number_format, cell.border, cell.font = NUM, BOX, Font(name=FONT)
            r += 1
        r += 1
    ws.column_dimensions["A"].width = 45
    for k in range(2, 12):
        ws.column_dimensions[get_column_letter(k)].width = 18


@st.cache_data(show_spinner=False)
def make_excel(df, corp_name, div_label):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "inhyeok"
    for sj in df["sj_nm"].unique():
        part = df[df["sj_nm"] == sj]
        first = part.iloc[0]
        cols = amount_cols(part)
        ws = wb.create_sheet(str(sj)[:31])
        ws.sheet_view.showGridLines = False
        if cols:
            ws.merge_cells(f"A1:{get_column_letter(len(cols) + 1)}1")
        ws["A1"] = f"{sj} ({div_label})"
        ws["A1"].font = Font(name=FONT, size=14, bold=True)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws["A2"] = f"{first.get('thstrm_nm', '')} : {first.get('thstrm_dt', '')}"
        ws["A3"] = f"{first.get('frmtrm_nm', '')} : {first.get('frmtrm_dt', '')}"
        ws["A4"] = corp_name
        ws["A5"], ws["A5"].font = AUTHOR, CREDIT
        if str(first.get("sj_div")) == "SCE":
            write_sce(ws, part, cols, 6)
            continue
        for c, h in enumerate(["과목"] + [h for h, _ in cols], 1):
            cell = ws.cell(row=6, column=c, value=h)
            cell.font = Font(name=FONT, bold=True)
            cell.fill, cell.border = HEAD, BOX
            cell.alignment = Alignment(horizontal="center")
        r = 7
        for _, row in part.iterrows():
            nm = str(row["account_nm"])
            major = nm.endswith(("총계", "합계")) or nm in MAJOR
            detail = row.get("account_detail")
            if isinstance(detail, str) and detail not in ("", "-"):
                nm = f"{nm} | {detail}"
            ws.cell(row=r, column=1, value=nm).alignment = Alignment(indent=0 if major else 1)
            for k, (_, col) in enumerate(cols, 2):
                ws.cell(row=r, column=k, value=to_num(row.get(col))).number_format = NUM
            for c in range(1, len(cols) + 2):
                cell = ws.cell(row=r, column=c)
                cell.border = BOX
                cell.font = Font(name=FONT, bold=major)
                if major:
                    cell.fill = SUB
            r += 1
        ws.column_dimensions["A"].width = 45
        for k in range(2, len(cols) + 2):
            ws.column_dimensions[get_column_letter(k)].width = 22
        ws.freeze_panes = "A7"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ===== 2. 감사보고서 재무제표 + 주석 엑셀 =====
def to_number(s, strict=False):
    if not isinstance(s, str):
        return s
    t = s.replace(" ", "").strip()
    if t in ("-", "－"):
        return 0.0
    if strict and not AMOUNT.fullmatch(t):
        return s
    neg = t.startswith("(") and t.endswith(")")
    t2 = t.strip("()").replace(",", "")
    if re.fullmatch(r"-?\d+(\.\d+)?", t2):
        n = float(t2)
        return -n if neg else n
    return s


def span_of(cell, attr):
    try:
        return max(1, int(cell.get(attr, "1") or 1))
    except ValueError:
        return 1


def head_rows(t):
    """주석 표의 머리글 마지막 행: 숫자가 처음 나오는 행 바로 위까지, 첫 행의 세로 병합 범위까지"""
    end = 0
    for i in range(len(t)):
        if any(isinstance(to_number(v, True), float) and str(v).strip() not in ("-", "－")
               for v in t.iloc[i, 1:]):
            end = max(i - 1, 0)
            break
    for r0, _, r1, _ in t.attrs.get("merges", []):
        if r0 == 0:
            end = max(end, r1)
    return min(end, len(t) - 1)


def read_items(text):
    """원문을 위에서부터 읽어 ('p', 한 줄 문장) 또는 ('table', 표)를 순서대로 돌려줌"""
    text = re.sub(r"<\?xml[^>]*\?>", "", text)
    text = re.sub(r"<(/?)T[EU](\s|>)", r"<\1TD\2", text)
    text = re.sub(r"<TITLE\b[^>]*>", "<P>", text)              # 섹션 제목도 문장으로 읽기
    text = re.sub(r"</TITLE\s*>", "</P>", text)
    text = re.sub(r"&cr;", "<BR/>", text, flags=re.I)    # DART 원문의 줄바꿈 표시(&cr;) → 줄바꿈
    doc = lh.document_fromstring(text)
    for br in doc.iter("br"):
        br.tail = "\n" + (br.tail or "")
    for sp in doc.iter("span"):
        if "B" in (sp.get("usermark") or ""):
            sp.tail = "\n" + (sp.tail or "")
    for pe in doc.iter("p", "div", "li"):                         # 문단 경계 = 줄바꿈 (표 칸 안의 여러 문단이 붙지 않게)
        pe.tail = "\n" + (pe.tail or "")
    items, pending = [], [""]

    def flush():
        if pending[0]:
            items.append(("p", pending[0]))
            pending[0] = ""

    def add_line(line):
        s = " ".join(line.split())
        if not s:
            return
        parts = MID_HEAD.split(s)
        if len(parts) > 1:                           # 한 줄에 붙은 다음 주석 제목은 따로 떼기
            for part in parts:
                add_line(part)
            return
        if pending[0]:                               # "1." 뒤에 제목이 오면 합치기
            s, pending[0] = f"{pending[0]} {s}", ""
        if NUM_ONLY.match(s):
            pending[0] = s
            return
        items.append(("p", s))

    def own_rows(tbl):
        """이 표 자신의 행만 (안에 들어 있는 다른 표의 행은 제외)"""
        return [tr for tr in tbl.iter("tr") if next(tr.iterancestors("table"), None) is tbl]

    def walk(node):
        """표 안에 표가 들어 있는 '틀' 칸: 글은 문장으로, 안쪽 표는 표로 순서대로 꺼냄"""
        buf = [node.text or ""]

        def out_text():
            for line in "".join(buf).split("\n"):
                add_line(line)
            buf.clear()

        for child in node:
            if not isinstance(child.tag, str):
                buf.append(child.tail or "")
                continue
            if child.tag == "table":
                out_text()
                handle_table(child)
            elif child.find(".//table") is not None:
                out_text()
                walk(child)
            else:
                buf.append(child.text_content())
            buf.append(child.tail or "")
        out_text()

    def handle_table(el):
        if el.find(".//table") is not None:          # 틀 역할의 바깥 표 → 칸마다 풀어서 읽기
            for tr in own_rows(el):
                for cell in tr:
                    if cell.tag in ("td", "th"):
                        walk(cell)
            return
        rows, merges, carry = [], [], {}             # carry: 열 번호 → 위 칸 rowspan으로 남은 행 수
        for tr in el.iter("tr"):
            row, c, r = [], 0, len(rows)
            for cell in tr:
                if cell.tag not in ("td", "th"):
                    continue
                while carry.get(c, 0) > 0:           # 위에서 내려온 병합 칸은 비워 두고 건너뜀
                    carry[c] -= 1
                    row.append("")
                    c += 1
                cs, rs = span_of(cell, "colspan"), span_of(cell, "rowspan")
                row.append(" ".join(cell.text_content().split()))
                row.extend([""] * (cs - 1))
                if rs > 1:
                    for k in range(c, c + cs):
                        carry[k] = rs - 1
                if cs > 1 or rs > 1:
                    merges.append((r, c, r + rs - 1, c + cs - 1))
                c += cs
            last = max((k for k, v in carry.items() if v > 0), default=-1)
            while c <= last:                         # 행 끝쪽의 병합 칸
                if carry.get(c, 0) > 0:
                    carry[c] -= 1
                row.append("")
                c += 1
            if row:
                rows.append(row)
        if not rows:
            return
        cells = [c for r in rows for c in r if c]
        joined = " ".join(cells)
        filled = sum(1 for r in rows if any(r))
        # 제목이 1행짜리 표(번호칸/제목칸 분리 포함)에 들어 있는 경우 → 문장으로 취급
        if filled == 1 and joined and len(joined) <= 100 and not AMOUNT.search(joined):
            add_line(joined)
            return
        # 1칸짜리 상자 표 → 줄 단위 문장으로 풀어줌
        if len(cells) == 1:
            for line in el.text_content().split("\n"):
                add_line(line)
            return
        flush()
        w = max(len(r) for r in rows)
        t = pd.DataFrame([r + [""] * (w - len(r)) for r in rows])
        t.attrs["merges"] = merges
        items.append(("table", t))

    for el in doc.iter("p", "table"):
        if any(a.tag == "table" for a in el.iterancestors()):
            continue
        if el.tag == "p":
            for line in el.text_content().split("\n"):
                add_line(line)
            continue
        handle_table(el)
    flush()
    return items


def classify(t):
    """표의 첫 열 계정 이름으로 재무제표 종류를 판별. 해당 없으면 None"""
    if len(t) < 4 or t.shape[1] < 2:
        return None
    first = "".join(t.iloc[:, 0]).replace(" ", "")
    head = "".join(t.iloc[:4].values.ravel()).replace(" ", "")
    if "자산총계" in first and ("부채총계" in first or "자본총계" in first):
        return "재무상태표"
    if re.search(r"영업활동(으로인한)?현금흐름", first) and "투자활동" in first:
        return "현금흐름표"
    if "기초" in first and ("자본금" in head or "이익잉여금" in head):
        return "자본변동표"
    if "총포괄" in first and ("당기순" in first or "매출" in first):
        return "포괄손익계산서"
    if ("매출액" in first or "영업수익" in first) and ("영업이익" in first or "당기순" in first):
        return "손익계산서"
    return None


def split_head(val, start):
    """'28. 금융수익과 금융비용(1) 당기와...'처럼 제목과 본문이 붙은 줄을 (제목, 본문)으로 나눔"""
    sub = SUB_START.search(val, start)
    if sub and sub.start() - start <= 40:
        cut = sub.start()
    elif len(val) > 20:
        word = BODY_WORD.search(val, start + 2)
        cut = word.start() if word and word.start() - start <= 40 else min(len(val), start + 30)
    else:
        return val, ""
    return val[:cut].strip(), val[cut:].strip()


def split_audit(items, tol=0, body=False, consolidated=False):
    """재무제표 표와 주석(번호별 묶음)으로 나눔.
    found 는 원문에 나온 순서대로 채워짐(dict 삽입 순서 유지).
    tol>0 이면 번호가 몇 개 건너뛰어도(제목 누락 대비) 주석 제목으로 인정.
    body=True(반기·분기 본문)이면 다음 섹션 제목에서 멈춤"""
    found, sections, cur, after = {}, [], None, []
    for typ, val in items:
        if typ == "p":
            if found and len(val) <= 100 and STOP.search(val):
                break
            if body and found and len(val) <= 100 and (
                    BODY_STOP.search(val) or (consolidated and SEP_BS.match(val))):
                break
            if SECTION_NOTE.search(val):                 # 섹션 제목은 주석 번호로 보지 않음
                continue
            m = NOTE_HEAD.match(val) if found else None
            nxt = len(sections) + 1
            num = int(m.group(1)) if m else 0
            dot = bool(m) and not m.group(0).rstrip().endswith(")")
            # 짧은 줄은 번호 범위 안이면 제목으로, 긴 줄(제목+본문이 붙은 경우)은 'N.' 형식이고 다음 번호일 때만
            if m and ((len(val) <= 100 and nxt <= num <= nxt + tol) or (dot and num == nxt)):
                title, rest = split_head(val, m.end())
                title = re.sub(r"^주\s*석\s*", "", title).rstrip(" :：")
                cur = {"title": title, "items": [("p", rest)] if rest else []}
                sections.append(cur)
            elif cur is not None:
                cur["items"].append(("p", val))
            elif found:
                after.append(("p", val))
            continue
        t = val
        if found and STOP.search("".join(t.iloc[:2].values.ravel())):
            break
        if cur is not None:
            cur["items"].append(("table", t))
            continue
        kind = classify(t)
        if kind and kind not in found:
            found[kind] = t
        elif found:
            after.append(("table", t))           # 같은 종류의 두 번째 표도 버리지 않음
    return found, sections, after


def cut_to_statements(items, consolidated):
    """'2-1. 연결 재무상태표' 같은 제목 위치부터 잘라냄. 못 찾으면 그대로 반환"""
    pat = r"^\d+\s*-\s*\d+\s*\.\s*" + (r"연결\s*" if consolidated else "") + r"재무상태표"
    for i, (typ, val) in enumerate(items):
        if typ == "p" and re.match(pat, val):
            return items[i:]
    return items


@st.cache_data(show_spinner="감사보고서 원문을 불러오는 중...")
def get_audit_tables(rcept_no, consolidated):
    try:
        r = requests.get("https://opendart.fss.or.kr/api/document.xml",
                         params={"crtfc_key": API_KEY, "rcept_no": rcept_no}, timeout=60)
    except requests.RequestException as e:
        return {}, [], f"원문 요청 실패: {safe_err(e)}"
    try:
        z = zipfile.ZipFile(io.BytesIO(r.content))
    except zipfile.BadZipFile:
        m = re.search(rb"<message>(.*?)</message>", r.content, re.S)
        detail = m.group(1).decode("utf-8", errors="ignore") if m else ""
        return {}, [], f"원문 파일을 받지 못했습니다. {detail}".strip()
    target, is_body = None, False
    for f in z.namelist():
        raw = z.read(f)
        m = re.search(rb'encoding="([^"]+)"', raw[:300])
        enc = m.group(1).decode() if m else "utf-8"
        if enc.lower().replace("-", "") in ("euckr", "ksc5601"):
            enc = "cp949"
        try:
            text = raw.decode(enc, errors="ignore")
        except LookupError:
            text = raw.decode("utf-8", errors="ignore")
        doc_name = re.search(r"<DOCUMENT-NAME[^>]*>(.*?)</DOCUMENT-NAME>", text, re.S)
        doc_name = doc_name.group(1).strip() if doc_name else ""
        is_audit = "감사보고서" in doc_name and ("연결" in doc_name) == consolidated
        # [기재정정] 등 접두어가 붙은 분기/반기보고서 본문도 인식
        is_body = bool(re.search(r"(분기|반기)보고서", doc_name)) and "감사" not in doc_name
        if is_audit or is_body:
            target = text
            break
    if target is None:
        return {}, [], "첨부된 감사보고서를 찾지 못했습니다."
    items = read_items(target)
    if is_body:
        items = cut_to_statements(items, consolidated)
    opt = {"body": is_body, "consolidated": consolidated}
    found, sections, after = split_audit(items, **opt)             # 1차: 엄격한 번호 순서
    loose = split_audit(items, tol=3, **opt)                       # 2차: 번호 누락 허용
    if len(loose[1]) > len(sections):                              # 중간 번호를 못 찾아 끊겼으면 2차 결과 사용
        found, sections, after = loose
    if not sections and after:                              # 3차: 통째로 한 시트에 담기
        sections = [{"title": "주석(전체)", "items": after}]
    return found, sections, None


def write_table(ws, t, start_row, hdr_end, note_cols, strict, col0=1, merges=()):
    for i in range(len(t)):
        is_head = i <= hdr_end
        name0 = str(t.iat[i, 0])
        major = bool(re.match(r"^\s*[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩIVX]+\s*[\.\s]", name0)) \
            or name0.replace(" ", "").endswith(("총계", "합계"))
        for j in range(t.shape[1]):
            v = t.iat[i, j]
            val = v if (is_head or j == 0 or j in note_cols) else to_number(v, strict)
            cell = ws.cell(row=start_row + i, column=col0 + j, value=None if val == "" else val)
            cell.border = BOX
            if isinstance(val, float):
                cell.number_format = NUM
            if is_head:
                cell.font = Font(name=FONT, bold=True)
                cell.fill = HEAD
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            else:
                cell.font = Font(name=FONT, bold=major)
    taken = set()
    for r0, c0, r1, c1 in merges:                    # 원문의 병합 칸 그대로
        r1, c1 = min(r1, len(t) - 1), min(c1, t.shape[1] - 1)
        area = {(a, b) for a in range(r0, r1 + 1) for b in range(c0, c1 + 1)}
        if len(area) > 1 and not (area & taken):     # 겹치는 병합은 엑셀 오류가 나므로 건너뜀
            taken |= area
            ws.merge_cells(start_row=start_row + r0, start_column=col0 + c0,
                           end_row=start_row + r1, end_column=col0 + c1)
    return start_row + len(t)


def safe_sheet_name(title, used):
    name = re.sub(r"[\[\]\:\*\?\/\\]", "", title)[:31].strip(" '") or "주석"   # 엑셀은 ' 로 시작·끝나는 시트 이름 불가
    base, k = name, 2
    while name in used:
        name = f"{base[:28].strip(" '")}_{k}"
        k += 1
    used.add(name)
    return name


def make_audit_excel(found, sections, corp_name, label):
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = "inhyeok"
    used = {"목차"}
    link = Font(name=FONT, color="0563C1", underline="single")

    toc = wb.create_sheet("목차")
    toc.sheet_view.showGridLines = False
    toc["B2"] = f"{corp_name} {label}"
    toc["B2"].font = Font(name=FONT, size=14, bold=True)
    toc["B3"], toc["B3"].font = AUTHOR, CREDIT
    toc.column_dimensions["A"].width = 2
    toc.column_dimensions["B"].width = 60
    toc_row = 4

    for kind in found:                                  # 원문에 나온 순서 그대로
        merges = found[kind].attrs.get("merges", [])
        t = found[kind].reset_index(drop=True)
        ws = wb.create_sheet(kind)
        used.add(kind)
        ws.sheet_view.showGridLines = False
        ws["B2"] = f"{kind} ({label})"
        ws["B2"].font = Font(name=FONT, size=14, bold=True)
        ws["B3"] = corp_name
        ws["B4"], ws["B4"].font = AUTHOR, CREDIT
        nospace = t.apply(lambda c: c.astype(str).str.replace(" ", ""))
        hdr_end = next((i for i in range(len(t)) if nospace.iloc[i].str.contains("과목").any()), 0)
        note_cols = {j for j in range(t.shape[1])
                     if nospace.iloc[:hdr_end + 1, j].str.contains("주석").any()}
        write_table(ws, t, 5, hdr_end, note_cols, strict=False, col0=2, merges=merges)
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 40
        for j in range(3, t.shape[1] + 2):
            ws.column_dimensions[get_column_letter(j)].width = 18
        ws.freeze_panes = ws.cell(row=6 + hdr_end, column=1)
        c = toc.cell(row=toc_row, column=2, value=kind)
        c.hyperlink, c.font = f"#'{kind}'!A1", link
        toc_row += 1

    toc_row += 1
    for sec in sections:
        sname = safe_sheet_name(sec["title"], used)
        ws = wb.create_sheet(sname)
        ws.sheet_view.showGridLines = False
        ws["B2"] = sec["title"]
        ws["B2"].font = Font(name=FONT, size=12, bold=True)
        ws["B3"] = corp_name
        ws["B3"].font = Font(name=FONT, color="808080")
        ws["B4"], ws["B4"].font = AUTHOR, CREDIT
        r, widest = 5, 1
        for typ, val in sec["items"]:
            if typ == "p":
                cell = ws.cell(row=r, column=2, value=val)
                cell.font = Font(name=FONT, bold=bool(SUBHEAD.match(val)))
                r += 2                                   # 문장 사이에 빈 줄 하나
            else:
                merges = val.attrs.get("merges", [])
                t = val.reset_index(drop=True)
                t.attrs["merges"] = merges
                r = write_table(ws, t, r, head_rows(t), set(), strict=True, col0=2, merges=merges) + 1
                widest = max(widest, t.shape[1])
        ws.column_dimensions["A"].width = 2
        ws.column_dimensions["B"].width = 30
        for j in range(3, widest + 2):
            ws.column_dimensions[get_column_letter(j)].width = 16
        c = toc.cell(row=toc_row, column=2, value=sec["title"])
        c.hyperlink, c.font = f"#'{sname}'!A1", link
        toc_row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ===== 3. 데이터 불러오기 =====
DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "dart.db")


def has_corp_table(con):
    return con.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='corp_code'").fetchone() is not None


def build_corp_db(con):
    """dart.db가 없으면 OpenDART 고유번호 파일(corpCode.xml)을 받아 회사 목록 테이블을 만듦"""
    r = requests.get("https://opendart.fss.or.kr/api/corpCode.xml",
                     params={"crtfc_key": API_KEY}, timeout=60)
    try:
        z = zipfile.ZipFile(io.BytesIO(r.content))
    except zipfile.BadZipFile:
        m = re.search(rb"<message>(.*?)</message>", r.content, re.S)
        raise RuntimeError("회사 목록을 받지 못했습니다. API 키를 확인하세요. "
                           + (m.group(1).decode("utf-8", errors="ignore") if m else ""))
    from lxml import etree
    root = etree.fromstring(z.read(z.namelist()[0]))
    rows = [((e.findtext("corp_code") or "").strip(), (e.findtext("corp_name") or "").strip(),
             (e.findtext("stock_code") or "").strip(), (e.findtext("modify_date") or "").strip())
            for e in root.iter("list")]
    con.execute("CREATE TABLE corp_code (corp_code TEXT, corp_name TEXT, stock_code TEXT, modify_date TEXT)")
    con.executemany("INSERT INTO corp_code VALUES (?, ?, ?, ?)", rows)
    con.commit()


@st.cache_data(show_spinner="회사 목록을 준비하는 중... (처음 한 번만, 1분 정도)")
def load_corp():
    con = sqlite3.connect(DB_PATH)
    try:
        if not has_corp_table(con):
            build_corp_db(con)
        df = pd.read_sql("SELECT corp_code, corp_name, stock_code, modify_date FROM corp_code "
                         "WHERE stock_code IS NOT NULL AND TRIM(stock_code) != ''", con)
    finally:
        con.close()
    return df


@st.cache_data(show_spinner="DART에서 불러오는 중...")
def get_fs(corp_code, year, reprt, div):
    try:
        res = requests.get(URL, params={"crtfc_key": API_KEY, "corp_code": corp_code,
                                        "bsns_year": year, "reprt_code": reprt, "fs_div": div},
                           timeout=30).json()
    except (requests.RequestException, ValueError) as e:
        return None, f"요청 실패: {safe_err(e)}"
    if res.get("status") != "000":
        return None, res.get("message")
    df = pd.DataFrame(res.get("list", []))
    if df.empty:
        return None, "데이터가 비어 있습니다."
    df["_s"] = df["sj_div"].map(SJ_RANK).fillna(9)
    df["_o"] = pd.to_numeric(df["ord"], errors="coerce")
    df = df.sort_values(["_s", "_o"], kind="stable").drop(columns=["_s", "_o"]).reset_index(drop=True)
    return df, None


# ===== 원문 순서로 XBRL 계정 재정렬 =====
ORIG_KIND = {"BS": ("재무상태표",), "IS": ("손익계산서", "포괄손익계산서"),
             "CIS": ("포괄손익계산서", "손익계산서"), "CF": ("현금흐름표",)}


# 회사마다 표기가 다른 같은 계정 → 하나로 통일
NM_ALIAS = [(r"^(자본과부채|부채와자본|부채및자본|자본및부채)총계$", "부채와자본총계"),
            (r"^(지배기업의?소유주에게귀속되는자본|지배기업소유주지분|지배기업소유지분)$", "지배기업소유주지분"),
            (r"^비지배(주주)?지분$", "비지배지분")]


def norm_nm(s):
    """계정명 비교용: 번호·괄호 내용(주석번호, '사채 제외' 등)·공백·기호 제거"""
    s = re.sub(r"\([^)]*\)|（[^）]*）", "", str(s))
    s = re.sub(r"^\s*[ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ]+\s*[\.\s]", "", s)
    s = re.sub(r"^\s*(\d{1,2}|[IVX]+|[가-하])\s*[\.\)]\s*", "", s)
    s = re.sub(r"[\s\.,·ㆍ\-\[\]]", "", s)
    for pat, rep in NM_ALIAS:
        s = re.sub(pat, rep, s)
    return s


def reorder_like_original(df, found, min_rate=0.5):
    """원문 재무제표 표의 행 순서에 맞춰 XBRL 행을 재정렬.
    1차로 계정명이 정확히 같은 행을 모두 배정한 뒤, 남은 행만 비슷한 이름(difflib)으로 매칭.
    매칭 안 된 행은 XBRL상 바로 앞 행 뒤에 붙임. 매칭률이 min_rate 미만이면 기존 순서 유지."""
    out, rates = [], {}
    for sj_div, part in df.groupby("sj_div", sort=False):
        t = next((found[k] for k in ORIG_KIND.get(sj_div, ()) if k in found), None)
        if t is None:
            out.append(part)
            continue
        index = {}
        for p, n in enumerate(norm_nm(x) for x in t.iloc[:, 0]):
            if n:
                index.setdefault(n, []).append(p)
        names = [norm_nm(x) for x in part["account_nm"]]
        pos, used = [None] * len(names), set()

        def take(i, key):
            p = next((p for p in index.get(key, []) if p not in used), None)
            if p is not None:
                pos[i] = p
                used.add(p)
            return p is not None

        for i, n in enumerate(names):                    # 1차: 정확히 일치
            if n:
                take(i, n)
        for i, n in enumerate(names):                    # 2차: 비슷한 이름
            if pos[i] is None and n:
                cands = [x for x, ps in index.items() if any(p not in used for p in ps)]
                m = difflib.get_close_matches(n, cands, n=1, cutoff=0.6)
                if m:
                    take(i, m[0])
        keys, last = [], -1.0
        for p in pos:                                    # 3차: 못 찾은 행은 앞 행 바로 뒤
            last = last + 1e-3 if p is None else float(p)
            keys.append(last)
        rates[sj_div] = sum(p is not None for p in pos) / len(pos)
        if rates[sj_div] < min_rate:
            out.append(part)
        else:
            out.append(part.assign(_k=keys).sort_values("_k", kind="stable").drop(columns="_k"))
    return pd.concat(out, ignore_index=True), rates


def value(df, label, basis="3개월"):
    ids = [k for k, v in IDS.items() if v == label]
    rows = df.loc[df["account_id"].isin(ids) & df["sj_div"].isin(SJ_OF[label])]
    order = ("thstrm_add_amount", "thstrm_amount") if basis == "누적" else ("thstrm_amount", "thstrm_add_amount")
    for col in order:
        if col in rows.columns:
            s = pd.to_numeric(rows[col], errors="coerce").dropna()
            if len(s):
                return float(s.iloc[0])
    return None


# ===== 5개년 재무분석 데이터 =====
# 항목: (찾을 재무제표, 표준 계정 ID, ID가 없을 때 쓸 계정명 패턴(공백 제거 후 비교))
FIVE = {
    "매출액": (("IS", "CIS"), ("ifrs-full_Revenue",), r"^(매출액|매출|영업수익|수익\(매출액\))$"),
    "영업이익": (("IS", "CIS"), ("dart_OperatingIncomeLoss", "ifrs-full_ProfitLossFromOperatingActivities"),
             r"^영업이익(\(손실\))?$"),
    "당기순이익": (("IS", "CIS"), ("ifrs-full_ProfitLoss",), r"^(연결)?당기순이익(\(손실\))?$"),
    "자산총계": (("BS",), ("ifrs-full_Assets",), r"^자산총계$"),
    "유동자산": (("BS",), ("ifrs-full_CurrentAssets",), r"^유동자산$"),
    "부채총계": (("BS",), ("ifrs-full_Liabilities",), r"^부채총계$"),
    "자본총계": (("BS",), ("ifrs-full_Equity",), r"^자본총계$"),
    "영업활동현금흐름": (("CF",), ("ifrs-full_CashFlowsFromUsedInOperatingActivities",), r"^영업활동.*현금흐름"),
    "유형자산취득": (("CF",), ("ifrs-full_PurchaseOfPropertyPlantAndEquipmentClassifiedAsInvestingActivities",),
               r"^유형자산의?취득"),
    "무형자산취득": (("CF",), ("ifrs-full_PurchaseOfIntangibleAssetsClassifiedAsInvestingActivities",),
               r"^무형자산의?취득"),
}


def pick_amount(df, key, cols=("thstrm_amount",)):
    """보고서 데이터에서 항목 금액 하나를 찾음 (ID 우선, 없으면 계정명). cols는 읽을 칸의 우선순위"""
    if df is None:
        return None
    sj, ids, pat = FIVE[key]
    part = df[df["sj_div"].isin(sj)]
    rows = part[part["account_id"].isin(ids)]
    if rows.empty:
        nm = part["account_nm"].astype(str).str.replace(r"\s", "", regex=True)
        rows = part[nm.str.match(pat)]
    for c in cols:
        if c in rows.columns:
            v = pd.to_numeric(rows[c], errors="coerce").dropna()
            if len(v):
                return float(v.iloc[0])
    return None


def sub(a, b):
    return None if a is None or b is None else a - b


IS_KEYS = ("매출액", "영업이익", "당기순이익")
BS_KEYS = ("자산총계", "유동자산", "부채총계", "자본총계")
CF_KEYS = ("영업활동현금흐름", "유형자산취득")
QTR = {1: "11013", 2: "11012", 3: "11014", 4: "11011"}     # 분기 → 보고서 코드 (4분기 = 사업보고서)
CUM = ("thstrm_add_amount", "thstrm_amount")                 # 누적 금액 (현금흐름표는 분기보고서도 누적)


def year_values(corp_code, y, div):
    fs, _ = get_fs(corp_code, str(y), "11011", div)
    if fs is None:
        return None
    rec = {"기간": str(y)}
    for k in IS_KEYS + BS_KEYS + CF_KEYS:
        rec[k] = pick_amount(fs, k)
    return rec


def quarter_values(corp_code, y, q, div):
    """한 분기의 3개월 금액. 4분기 손익 = 연간 − 3분기 누적, 현금흐름 = 이번 누적 − 직전 분기 누적"""
    fs, _ = get_fs(corp_code, str(y), QTR[q], div)
    if fs is None:
        return None
    before = get_fs(corp_code, str(y), QTR[q - 1], div)[0] if q > 1 else None
    rec = {"기간": f"{y}/{q * 3:02d}"}
    for k in IS_KEYS:
        if q < 4:
            rec[k] = pick_amount(fs, k)                                          # 3개월
            rec[k + "_전년"] = pick_amount(fs, k, ("frmtrm_q_amount", "frmtrm_amount"))
        else:
            rec[k] = sub(pick_amount(fs, k), pick_amount(before, k, CUM))
            rec[k + "_전년"] = sub(pick_amount(fs, k, ("frmtrm_amount",)), pick_amount(before, k, ("frmtrm_add_amount",)))
    for k in BS_KEYS:
        rec[k] = pick_amount(fs, k)
    for k in CF_KEYS:
        cur = pick_amount(fs, k, CUM)
        rec[k] = cur if q == 1 else sub(cur, pick_amount(before, k, CUM))
    return rec


@st.cache_data(show_spinner="재무분석 데이터를 불러오는 중...")
def analysis_data(corp_code, end_year, div, mode):
    """최근 5개 기간(연간: 사업보고서 5개년 / 분기: 최근 5개 분기)을 모아 지표 계산 (억원, %)"""
    recs, y = [], int(end_year)
    if mode == "연간":
        while len(recs) < 5 and y >= max(2015, int(end_year) - 7):      # OpenDART는 2015년부터 제공
            r = year_values(corp_code, y, div)
            if r:
                recs.append(r)
            y -= 1
    else:
        q, tries = 4, 0
        while len(recs) < 5 and tries < 10 and y >= 2015:
            r = quarter_values(corp_code, y, q, div)
            if r:
                recs.append(r)
            tries += 1
            q -= 1
            if q == 0:
                q, y = 4, y - 1
    t = pd.DataFrame(recs[::-1])
    if t.empty:
        return t
    t = t.set_index("기간").astype(float) / 1e8
    t["CAPEX"] = t["유형자산취득"].abs()
    t["잉여현금흐름"] = t["영업활동현금흐름"] - t["CAPEX"].fillna(0)
    t["영업이익률"] = t["영업이익"] / t["매출액"] * 100
    t["순이익률"] = t["당기순이익"] / t["매출액"] * 100
    t["부채비율"] = t["부채총계"] / t["자본총계"] * 100

    def growth(cur, prev):
        return (cur - prev) / prev.abs() * 100                # 전기가 음수여도 방향이 맞도록 절댓값으로 나눔

    for k, g in (("매출액", "매출액증가율"), ("영업이익", "영업이익증가율"), ("당기순이익", "순이익증가율")):
        # 분기: 전년 같은 분기 대비 / 연간: 전년 대비
        t[g] = growth(t[k], t[k + "_전년"]) if mode == "분기" else growth(t[k], t[k].shift(1))
    for k, g in (("자산총계", "총자산증가율"), ("유동자산", "유동자산증가율"),
                 ("부채총계", "부채증가율"), ("자본총계", "자본증가율")):
        t[g] = growth(t[k], t[k].shift(1))                     # 분기: 직전 분기 말 대비
    return t.reset_index()


BAR_COLORS = ["#1f63d6", "#c8320f", "#7cbf1e"]
LINE_COLORS = ["#a070dc", "#f59c1a", "#2bb5a6", "#666666"]


def combo_chart(t, bars, lines=(), left="억원", right="%"):
    """막대(왼쪽 축) + 꺾은선(오른쪽 축) 차트"""
    x = alt.X("기간:N", title=None, axis=alt.Axis(labelAngle=0))
    lb = t.melt("기간", list(bars), "항목", "값")
    layers = [alt.Chart(lb).mark_bar().encode(
        x=x, xOffset=alt.XOffset("항목:N", sort=list(bars)),
        y=alt.Y("값:Q", title=f"[{left}]"),
        color=alt.Color("항목:N", sort=list(bars), title=None,
                        scale=alt.Scale(domain=list(bars), range=BAR_COLORS[:len(bars)]),
                        legend=alt.Legend(orient="bottom")),
        tooltip=["기간", "항목", alt.Tooltip("값:Q", format=",.0f")])]
    if lines:
        ll = t.melt("기간", list(lines), "항목", "값").dropna()
        layers.append(alt.Chart(ll).mark_line(point=True).encode(
            x=x, y=alt.Y("값:Q", title=f"[{right}]"),
            color=alt.Color("항목:N", sort=list(lines), title=None,
                            scale=alt.Scale(domain=list(lines), range=LINE_COLORS[:len(lines)]),
                            legend=alt.Legend(orient="bottom")),
            tooltip=["기간", "항목", alt.Tooltip("값:Q", format=",.1f")]))
    return alt.layer(*layers).resolve_scale(y="independent", color="independent").properties(height=330)


def title_tip(box, text, tip):
    """그래프 제목 옆에 (?) 표시 — 커서를 올리면 계산 근거가 보임"""
    tip = html.escape(tip).replace("\n", "&#10;")
    box.markdown(f'**{html.escape(text)}** <span title="{tip}" style="cursor:help; display:inline-block; '
                 'width:1.3em; height:1.3em; line-height:1.2em; text-align:center; border:1px solid #999; '
                 'border-radius:50%; font-size:0.75em; color:#666;">?</span>', unsafe_allow_html=True)


FORMULA = r"""| 분류 | 항목 | 계산 근거 |
|---|---|---|
| 손익 | 영업이익률 | 영업이익 ÷ 매출액 × 100 |
| 손익 | 순이익률 | 당기순이익 ÷ 매출액 × 100 |
| 손익 | 증가율 | (당기 − 비교기간) ÷ \|비교기간\| × 100 |
| 재무상태 | 부채비율 | 부채총계 ÷ 자본총계 × 100 |
| 현금흐름 | CAPEX | 유형자산의 취득 |
| 현금흐름 | 잉여현금흐름 | 영업활동현금흐름 − CAPEX |
| 분기 | 4분기 손익 | 사업보고서 연간 − 3분기보고서 누적 |
| 분기 | 분기 현금흐름 | 이번 분기 누적 − 직전 분기 누적 |

* 비교기간: 연간은 전년, 분기 손익은 전년 같은 분기, 분기 재무상태는 직전 분기 말
* 비교기간 값이 음수여도 증감 방향이 맞도록 절댓값으로 나눕니다.
* 출처: OpenDART 단일회사 전체 재무제표(XBRL)"""


def growth_chart(t, cols, skip_first=True):
    """증가율 꺾은선 차트 (skip_first: 첫 기간은 비교 대상이 없어 제외)"""
    lg = (t.iloc[1:] if skip_first else t).melt("기간", list(cols), "항목", "값").dropna()
    line = alt.Chart(lg).mark_line(point=True).encode(
        x=alt.X("기간:N", title=None, axis=alt.Axis(labelAngle=0)),
        y=alt.Y("값:Q", title="[%]"),
        color=alt.Color("항목:N", sort=list(cols), title=None,
                        scale=alt.Scale(domain=list(cols), range=LINE_COLORS[:len(cols)]),
                        legend=alt.Legend(orient="bottom")),
        tooltip=["기간", "항목", alt.Tooltip("값:Q", format=",.1f")])
    zero = alt.Chart(pd.DataFrame({"y": [0]})).mark_rule(color="#bbbbbb").encode(y="y:Q")
    return (zero + line).properties(height=330)


# ===== 4. 화면 =====
try:
    corp = load_corp()
except Exception as e:
    st.error(f"회사 목록을 불러오지 못했습니다: {safe_err(e)}")
    st.stop()
q = st.text_input("회사 이름을 입력하세요", "오리온")
q = q.strip()
cand = corp[corp["corp_name"].str.contains(q, case=False, na=False, regex=False)].copy()
nm, ql = cand["corp_name"].str.lower(), q.lower()
cand["_r"] = (nm != ql).astype(int) + (~nm.str.startswith(ql)).astype(int)  # 0: 정확히 일치, 1: 시작, 2: 포함
cand["_l"] = nm.str.len()
cand = cand.sort_values(["_r", "_l", "modify_date"], ascending=[True, True, False]) \
           .drop(columns=["_r", "_l"]).reset_index(drop=True)
if cand.empty:
    st.warning("검색 결과가 없습니다.")
    st.stop()

labels = cand["corp_name"] + " (" + cand["stock_code"] + ")"
dup = cand["corp_name"].duplicated(keep=False)             # 같은 이름이 여러 개면 갱신 연도 표시
labels[dup] = labels[dup] + " · 정보 갱신 " + cand.loc[dup, "modify_date"].str[:4] + "년"
i = st.selectbox("회사 선택", range(len(cand)), format_func=lambda k: labels[k])
code, name = cand.loc[i, "corp_code"], cand.loc[i, "corp_name"]

c1, c2, c3 = st.columns(3)
year = c1.selectbox("사업연도", [str(y) for y in range(date.today().year, 2014, -1)], index=1)
reprt = c2.selectbox("보고서", list(REPRT))
div = c3.radio("기준", list(DIV), horizontal=True)

df, err = get_fs(code, year, REPRT[reprt], DIV[div])
if df is None:
    st.error(f"자료가 없습니다: {err}")
    st.stop()

SJ_KO = {"BS": "재무상태표", "IS": "손익계산서", "CIS": "포괄손익계산서", "CF": "현금흐름표"}
if st.checkbox("원문(감사보고서) 순서로 계정 정렬", value=True,
               help="DART XBRL 데이터의 계정 순서가 원문과 다를 때 원문 표 순서에 맞춥니다."):
    found0, _, msg0 = get_audit_tables(str(df["rcept_no"].iloc[0]), div == "연결")
    if found0:
        df, rates = reorder_like_original(df, found0)
        low = [SJ_KO.get(k, k) for k, v in rates.items() if v < 0.5]
        if low:
            st.caption(f"원문과 계정명이 많이 달라 DART 제공 순서를 유지한 표: {', '.join(low)}")
    else:
        st.caption(f"원문을 찾지 못해 DART 제공 순서로 표시합니다. {msg0 or ''}")

basis = "3개월"
if reprt != "사업보고서":
    basis = st.radio("요약 숫자 기준 (손익 항목)", ["3개월", "누적"], horizontal=True,
                     help="3개월: 해당 분기만 / 누적: 1월부터 해당 분기 말까지")

rev = value(df, "매출액", basis)
op = value(df, "영업이익", basis)
liab = value(df, "부채총계", basis)
eq = value(df, "자본총계", basis)
tag = f" ({basis})" if reprt != "사업보고서" else ""


def eok(v):
    return f"{v / 1e8:,.0f}억" if v is not None else "-"


m1, m2, m3, m4 = st.columns(4)
m1.metric(f"매출액{tag}", eok(rev))
m2.metric(f"영업이익{tag}", eok(op))
m3.metric("부채비율", f"{liab / eq * 100:,.1f}%" if liab is not None and eq else "-")
m4.metric(f"영업이익률{tag}", f"{op / rev * 100:,.1f}%" if op is not None and rev else "-")

sheets = list(df["sj_nm"].unique())
for tab, sj in zip(st.tabs(sheets), sheets):
    part = df[df["sj_nm"] == sj]
    cols = amount_cols(part)
    if not cols:
        tab.info("표시할 금액이 없습니다.")
        continue
    if str(part["sj_div"].iloc[0]) == "SCE":
        pick = tab.radio("기간", [h for h, _ in cols], horizontal=True, key=f"sce_{sj}")
        f = sce_frame(part, dict(cols)[pick]) / 1e8
        f.index.name = "계정 (억원)"
        tab.dataframe(f.style.format("{:,.1f}", na_rep="-"), use_container_width=True)
        continue
    t = part[["account_nm"] + [c for _, c in cols]].copy()
    for _, c in cols:
        t[c] = pd.to_numeric(t[c], errors="coerce") / 1e8
    t.columns = ["계정"] + [f"{h}(억원)" for h, _ in cols]
    tab.dataframe(t.style.format({c: "{:,.1f}" for c in t.columns[1:]}, na_rep="-"),
                  use_container_width=True, hide_index=True)

d1, d2 = st.columns(2)
d1.download_button("📥 엑셀 파일로 다운로드 (서식 포함)", make_excel(df, name, div),
                   file_name=f"{name}_{year}_{reprt}_{div}.xlsx",
                   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
d2.download_button("CSV 다운로드 (원본 데이터)", df.to_csv(index=False).encode("utf-8-sig"),
                   file_name=f"{name}_{year}_{reprt}_{div}.csv", mime="text/csv")

# ===== 재무분석 그래프 (연간 5개년 / 최근 5개 분기) =====
st.divider()
h1, h2, h3, h4 = st.columns([3, 1.2, 1.2, 0.6])
h1.subheader("📈 재무분석")
fdiv = h2.selectbox("기준", list(DIV), index=list(DIV).index(div), key="five_div",
                    format_func=lambda d: f"K-IFRS({d})")
mode = h3.radio("기간", ["연간", "분기"], horizontal=True, key="five_mode")
with h4.popover("산식 ?"):
    st.markdown(FORMULA)
five = analysis_data(code, year, DIV[fdiv], mode)
if len(five) < 2:
    st.info("비교할 보고서가 2개 기간 이상 없어 그래프를 그릴 수 없습니다.")
else:
    is_q = mode == "분기"
    unit = "분기 3개월 기준 (4분기 = 연간 − 3분기 누적)" if is_q else "사업보고서 기준"
    st.caption(f"{five['기간'].iloc[0]}~{five['기간'].iloc[-1]} · {unit} · 단위: 억원, %")
    g1, g2, g3 = st.tabs(["포괄손익계산서", "재무상태표", "현금흐름표"])
    a, b = g1.columns(2)
    title_tip(a, "주요재무항목",
              ("매출액·영업이익·당기순이익: 분기별 3개월 금액 (4분기 = 연간 − 3분기 누적)\n" if is_q else
               "매출액·영업이익·당기순이익: 각 연도 사업보고서 금액\n")
              + "영업이익률 = 영업이익 ÷ 매출액 × 100\n순이익률 = 당기순이익 ÷ 매출액 × 100")
    a.altair_chart(combo_chart(five, ["매출액", "영업이익", "당기순이익"], ["영업이익률", "순이익률"]),
                   use_container_width=True)
    title_tip(b, "수익성장성지표",
              ("증가율 = (이번 분기 − 전년 같은 분기) ÷ |전년 같은 분기| × 100" if is_q else
               "증가율 = (당기 − 전기) ÷ |전기| × 100")
              + "\n대상: 매출액, 영업이익, 당기순이익")
    b.altair_chart(growth_chart(five, ["매출액증가율", "영업이익증가율", "순이익증가율"], skip_first=not is_q),
                   use_container_width=True)
    a, b = g2.columns(2)
    title_tip(a, "주요재무항목",
              ("자산총계·부채총계: 각 분기 말 잔액\n" if is_q else "자산총계·부채총계: 각 연도 말 잔액\n")
              + "부채비율 = 부채총계 ÷ 자본총계 × 100")
    a.altair_chart(combo_chart(five, ["자산총계", "부채총계"], ["부채비율"]), use_container_width=True)
    title_tip(b, "자산성장성지표",
              ("증가율 = (이번 분기 말 − 직전 분기 말) ÷ |직전 분기 말| × 100" if is_q else
               "증가율 = (당기말 − 전기말) ÷ |전기말| × 100")
              + "\n대상: 자산총계, 유동자산, 부채총계, 자본총계")
    b.altair_chart(growth_chart(five, ["총자산증가율", "유동자산증가율", "부채증가율", "자본증가율"]),
                   use_container_width=True)
    a, b = g3.columns(2)
    title_tip(a, "영업활동현금흐름 & CAPEX",
              ("분기 금액 = 이번 분기 누적 − 직전 분기 누적 (현금흐름표는 누적으로 공시)\n" if is_q else "")
              + "CAPEX = 유형자산의 취득 (현금흐름표 투자활동)")
    a.altair_chart(combo_chart(five, ["영업활동현금흐름", "CAPEX", "당기순이익"]), use_container_width=True)
    title_tip(b, "잉여현금흐름", "잉여현금흐름 = 영업활동현금흐름 − CAPEX\nCAPEX = 유형자산의 취득")
    b.altair_chart(combo_chart(five, ["잉여현금흐름"]), use_container_width=True)
    with st.expander("숫자로 보기"):
        show = five.set_index("기간")[["매출액", "영업이익", "당기순이익", "영업이익률", "순이익률",
                                      "자산총계", "부채총계", "자본총계", "부채비율",
                                      "영업활동현금흐름", "CAPEX", "잉여현금흐름"]].T
        st.dataframe(show.style.format("{:,.1f}", na_rep="-"), use_container_width=True)

# ===== 5. 감사보고서 양식 다운로드 =====
st.divider()
st.subheader("📑 원문 재무제표 + 주석 (회사 원본 양식)")
key = f"audit_{code}_{year}_{reprt}_{div}"
if st.button("원문 재무제표 불러오기"):
    found, sections, msg = get_audit_tables(str(df["rcept_no"].iloc[0]), div == "연결")
    if not found:
        st.error(msg or "원문에서 재무제표 표를 찾지 못했습니다.")
    else:
        st.session_state[key] = make_audit_excel(found, sections, name, f"{div} {reprt}")
        st.session_state[key + "_t"] = found
        st.success(f"찾은 표: {', '.join(found)} / 주석 {len(sections)}개")
if key in st.session_state:
    st.download_button("📥 원문 양식 엑셀 다운로드", st.session_state[key],
                       file_name=f"{name}_{year}_{reprt}_{div}_원문.xlsx",
                       mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    found = st.session_state[key + "_t"]
    kinds = list(found)                                  # 원문에 나온 순서 그대로
    for tab, k in zip(st.tabs(kinds), kinds):
        tab.dataframe(found[k], use_container_width=True, hide_index=True)